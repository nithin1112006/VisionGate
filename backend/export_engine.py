"""
Attenda Export Engine
Generates high-fidelity, production-grade Excel (.xlsx), PDF (.pdf), and CSV exports:
- Pure Python ReportLab integration for professional PDF tables and header styling
- OpenPyXL integration for styled multi-column workbooks with auto-fitted columns
- Standard CSV fallback / stream generator
- Conforms to Hallmark standards: no italic headings, clear typographic hierarchy
"""

import io
import csv
from datetime import datetime
from typing import List, Dict, Any, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch


def generate_csv(headers: List[str], rows: List[List[Any]]) -> str:
    """Generate a standard RFC 4180 CSV string."""
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([str(item) if item is not None else "" for item in row])
    return output.getvalue()


def generate_excel(
    title: str,
    sheet_name: str,
    headers: List[str],
    rows: List[List[Any]],
    subtitle: Optional[str] = None
) -> bytes:
    """Generate a beautifully formatted Excel workbook in bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Report")[:30]

    # Styles
    title_font = Font(name="Segoe UI", size=15, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=10, color="64748B")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    row_font = Font(name="Segoe UI", size=10, color="1F2937")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    current_row = 1

    # Title Banner
    ws.cell(row=current_row, column=1, value=title).font = title_font
    current_row += 1

    if subtitle:
        ws.cell(row=current_row, column=1, value=subtitle).font = subtitle_font
        current_row += 1

    # Export metadata timestamp
    meta_text = f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Attenda Institutional Portal"
    ws.cell(row=current_row, column=1, value=meta_text).font = subtitle_font
    current_row += 2  # extra spacing

    # Table Header
    header_row_idx = current_row
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if "date" in header.lower() or "status" in header.lower() or "count" in header.lower() or "period" in header.lower() else "left", vertical="center")
        cell.border = thin_border
    current_row += 1

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        is_even = (r_idx % 2 == 0)
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val if val is not None else "—")
            cell.font = row_font
            if is_even:
                cell.fill = alt_fill
            cell.border = thin_border
            # Format status badges
            val_str = str(val or "").strip()
            if val_str.lower() in ["present", "approved", "active", "success"]:
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="047857")
            elif val_str.lower() in ["absent", "rejected", "locked", "failed"]:
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")
            elif val_str.lower() in ["pending", "half day", "on duty", "warning"]:
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="B45309")
        current_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip title row from width calculation
            if cell.row < header_row_idx:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf_table(
    title: str,
    headers: List[str],
    rows: List[List[Any]],
    subtitle: Optional[str] = None,
    landscape_mode: bool = True
) -> bytes:
    """Generate a clean, print-ready PDF document with tabular data."""
    buf = io.BytesIO()
    page_size = landscape(A4) if landscape_mode else A4
    doc = SimpleDocTemplate(
        buf,
        pagesize=page_size,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    # Custom Roman Typography styles (no italics)
    title_style = ParagraphStyle(
        name="AttendaTitle",
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#0F172A"),
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        name="AttendaSubtitle",
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=12,
    )
    cell_style = ParagraphStyle(
        name="AttendaCell",
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
    )
    header_cell_style = ParagraphStyle(
        name="AttendaHeaderCell",
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white,
    )

    story = []

    # Title & Subtitle
    story.append(Paragraph(f"<b>{title}</b>", title_style))
    meta_subtitle = f"{subtitle + ' — ' if subtitle else ''}Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')} | Attenda Institutional Platform"
    story.append(Paragraph(meta_subtitle, subtitle_style))
    story.append(Spacer(1, 8))

    # Wrap cells in Paragraph to support auto-wrapping
    table_data = []
    header_row = [Paragraph(f"<b>{h}</b>", header_cell_style) for h in headers]
    table_data.append(header_row)

    for row in rows:
        row_cells = []
        for item in row:
            val = str(item) if item is not None else "—"
            row_cells.append(Paragraph(val, cell_style))
        table_data.append(row_cells)

    # Compute reasonable column widths
    avail_width = doc.width
    num_cols = len(headers)
    col_width = avail_width / max(num_cols, 1)

    t = Table(table_data, colWidths=[col_width] * num_cols, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
    ]))

    story.append(t)
    doc.build(story)
    return buf.getvalue()
